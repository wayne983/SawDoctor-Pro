import importlib.util
import sys
import unittest
from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory

import openpyxl


PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

spec = importlib.util.spec_from_file_location("sawdoctor_app", PROJECT_ROOT / "鋸片醫生.py")
app = importlib.util.module_from_spec(spec)
assert spec.loader is not None
spec.loader.exec_module(app)

from quote_domain import QuoteLineDraft


def batch(customer):
    return app.QuoteBatch(
        id="Q-260812-001",
        customer=customer,
        month_range="2026-08",
        created_at=datetime(2026, 8, 12, 9, 30).isoformat(),
        created_by="測試人員",
        price_version_id="PV-20260813",
        subtotal=0,
        output_file="",
    )


def line(brand_id, subtotal, *, grinding_unit=240):
    fanban_unit = subtotal - grinding_unit
    return QuoteLineDraft(
        blade_id=int(brand_id),
        customer="甲",
        brand_id=brand_id,
        od="305",
        thickness="3.0",
        teeth="100",
        spec="305 x 3.0 x 100T",
        source_grind="是",
        source_supp_teeth="2",
        source_fanban="是",
        grinding_qty=1,
        grinding_unit=grinding_unit,
        tooth_qty=0,
        tooth_unit=0,
        fanban_qty=1 if fanban_unit else 0,
        fanban_unit=fanban_unit,
        note=f"subtotal={subtotal}",
    )


class QuoteExcelTests(unittest.TestCase):
    def test_export_keeps_each_blade_as_a_separate_row_and_writes_total(self):
        with TemporaryDirectory() as directory:
            output = app.export_customer_quote_excel(
                Path(directory), batch("甲"), [line("4012", 540), line("4015", 240)]
            )
            ws = openpyxl.load_workbook(output, data_only=True)["客戶維修明細"]
            values = [[cell.value for cell in row] for row in ws.iter_rows()]

            self.assertTrue(any(row[-1] == "4012" for row in values))
            self.assertTrue(any(row[-1] == "4015" for row in values))
            total_row = next(row for row in values if row[0] == "未稅總計")
            self.assertEqual(total_row[6], 780)

    def test_export_uses_actual_overridden_units_in_group_header(self):
        with TemporaryDirectory() as directory:
            changed = line("4012", 600, grinding_unit=300)
            ws = openpyxl.load_workbook(
                app.export_customer_quote_excel(Path(directory), batch("甲"), [changed]),
                data_only=True,
            ).active

            self.assertTrue(
                any(cell.value == "$300" for row in ws.iter_rows() for cell in row)
            )

    def test_export_lists_all_actual_units_when_a_spec_has_overrides(self):
        with TemporaryDirectory() as directory:
            ws = openpyxl.load_workbook(
                app.export_customer_quote_excel(
                    Path(directory),
                    batch("甲"),
                    [line("4012", 540), line("4015", 600, grinding_unit=300)],
                ),
                data_only=True,
            ).active

            self.assertTrue(
                any(cell.value == "$240 / $300" for row in ws.iter_rows() for cell in row)
            )


if __name__ == "__main__":
    unittest.main()
