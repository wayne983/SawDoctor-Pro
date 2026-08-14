import importlib.util
import sys
import unittest
import xml.etree.ElementTree as ET
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile

import openpyxl

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

spec = importlib.util.spec_from_file_location("sawdoctor_dispatch_app", PROJECT_ROOT / "鋸片醫生.py")
app = importlib.util.module_from_spec(spec)
assert spec.loader is not None
spec.loader.exec_module(app)


def blade(customer, brand_id, od, thickness, teeth, *, supp_teeth="-", supp_seats="-",
          fanban="-", steel="-", receipt_date="2026-08-12", late_deadline="2026-08-28"):
    return {
        "customer": customer,
        "brand_id": brand_id,
        "od": od,
        "thickness": thickness,
        "teeth": teeth,
        "supp_teeth": supp_teeth,
        "supp_seats": supp_seats,
        "fanban": fanban,
        "steel": steel,
        "grind": "是",
        "folder_path": "",
        "receipt_date": receipt_date,
        "late_deadline": late_deadline,
        "ok_date": "",
        "status": "in_progress",
    }


class DispatchGroupingTests(unittest.TestCase):
    def test_group_for_export_combines_same_customer_and_spec_and_sums_repairs(self):
        items = app._group_for_export([
            blade("建通", "4001", "305", "3.0", "100", supp_teeth="2", supp_seats="1"),
            blade("建通", "4002", "305", "3.0", "100", supp_teeth="3", supp_seats="2", fanban="是", steel="是"),
        ])

        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["count"], 2)
        self.assertEqual(items[0]["spec"], "305/3.0T/100T")
        self.assertEqual(items[0]["repair_display"], "補齒x5、補座x3、反板x1、鋼面x1")
        self.assertEqual(items[0]["receipt_display"], "08/12")
        self.assertEqual(items[0]["deadline_display"], "08/28")

    def test_group_for_export_does_not_merge_different_customers_or_show_empty_repairs(self):
        items = app._group_for_export([
            blade("衛全", "7017", "195", "2.0", "80"),
            blade("建通", "7049", "195", "2.0", "80"),
        ])

        self.assertEqual([item["count"] for item in items], [1, 1])
        self.assertEqual([item["repair_display"] for item in items], ["", ""])


EXPECTED_HEADERS = [
    "勾選", "項次", "進貨日", "客戶", "規格",
    "數量", "維修項目", "備註", "交期",
]
WORD_NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def read_word_table(path):
    with ZipFile(path) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    table = root.find(".//w:tbl", WORD_NS)
    rows = table.findall("w:tr", WORD_NS)
    cells = [row.findall("w:tc", WORD_NS) for row in rows]
    text_rows = [
        ["".join(node.text or "" for node in cell.findall(".//w:t", WORD_NS)) for cell in row]
        for row in cells
    ]
    return text_rows, cells


class DispatchFormatTests(unittest.TestCase):
    def setUp(self):
        self.directory = TemporaryDirectory()
        self.original_base = app.BASE
        self.original_html = app.DISPATCH_HTML
        app.BASE = Path(self.directory.name)
        app.DISPATCH_HTML = app.BASE / "派工清單.html"
        self.rows = [
            blade("衛全", "7017", "195", "2.0", "80"),
            blade("衛全", "7049", "195", "2.0", "80"),
            blade("建通", "4001", "305", "3.0", "100", supp_teeth="2"),
        ]

    def tearDown(self):
        app.BASE = self.original_base
        app.DISPATCH_HTML = self.original_html
        self.directory.cleanup()

    def test_excel_uses_grouped_nine_columns_short_dates_and_blank_notes(self):
        ws = openpyxl.load_workbook(app.export_dispatch_excel(self.rows, "測試")).active

        self.assertEqual([cell.value for cell in ws[5]], EXPECTED_HEADERS)
        self.assertEqual(ws.cell(6, 3).value, "08/12")
        self.assertEqual(ws.cell(6, 4).value, "衛全")
        self.assertEqual(ws.cell(6, 5).value, "195/2.0T/80T")
        self.assertEqual(ws.cell(6, 6).value, 2)
        self.assertIsNone(ws.cell(6, 7).value)
        self.assertIsNone(ws.cell(6, 8).value)
        self.assertEqual(ws.cell(6, 9).value, "08/28")
        self.assertEqual(ws.cell(6, 1).font.sz, 14)
        self.assertNotIn("編號", [cell.value for cell in ws[5]])
        self.assertFalse(any("—" in str(cell.value or "") for row in ws.iter_rows(min_row=6, max_row=7, max_col=9) for cell in row))

    def test_html_uses_grouped_nine_columns_without_identifier(self):
        html = app.gen_dispatch_html(self.rows, "測試").read_text(encoding="utf-8")

        self.assertIn("<th>維修項目</th>", html)
        self.assertIn("<th>備註</th>", html)
        self.assertIn("195/2.0T/80T", html)
        self.assertIn("08/12", html)
        self.assertIn("08/28", html)
        self.assertNotIn("<th>編號</th>", html)
        self.assertNotIn("<td>—</td>", html)

    def test_word_uses_grouped_nine_columns_short_dates_and_blank_notes(self):
        output = app.export_dispatch_word(self.rows, "測試")
        rows, cells = read_word_table(output)
        with ZipFile(output) as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8")

        self.assertEqual(rows[0], EXPECTED_HEADERS)
        self.assertEqual(rows[1][2], "08/12")
        self.assertEqual(rows[1][4], "195/2.0T/80T")
        self.assertEqual(rows[1][5], "2")
        self.assertEqual(rows[1][7], "")
        self.assertEqual(rows[1][8], "08/28")
        self.assertTrue(any(node.get("{%s}val" % WORD_NS["w"]) == "28" for node in cells[1][0].findall(".//w:sz", WORD_NS)))
        self.assertNotIn("編號", document_xml)
        self.assertFalse(any("—" in text for row in rows for text in row))


if __name__ == "__main__":
    unittest.main()
